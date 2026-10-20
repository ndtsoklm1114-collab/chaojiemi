"""批註、頁首、文件屬性、xlsx 日期型生日欄、PDF 屬性、圖片 EXIF。全部用假資料。"""
import sys, pathlib, tempfile, subprocess, json, datetime
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
ROOT = pathlib.Path(__file__).resolve().parents[1]
PY = sys.executable


def deid(src, names, d):
    subprocess.run([PY, str(ROOT / "deidentify.py"), str(src), "--names", str(names), "--out", str(d / "out"), "--map", str(d / "map")], check=True, capture_output=True)
    return d / "out" / src.name, d / "map" / (src.name + ".對照表.json")


def run():
    ok = True
    with tempfile.TemporaryDirectory() as d:
        d = pathlib.Path(d)
        names = d / "names.txt"; names.write_text("[個案]\n王小明\n", encoding="utf-8")
        # docx：正文、頁首、批註、文件屬性
        import docx
        doc = docx.Document()
        doc.core_properties.author = "王小明"; doc.core_properties.title = "王小明 A123456789"
        doc.sections[0].header.paragraphs[0].text = "頁首 王小明"
        p = doc.add_paragraph(); r = p.add_run("正文 王小明 0912-345-678")
        doc.add_comment(r, text="批註裡的 王小明 B223456789", author="王小明")
        doc.save(d / "T.docx")
        out, mp = deid(d / "T.docx", names, d)
        o = docx.Document(out)
        from docx.oxml.ns import qn
        texts = []
        for part in o.part.package.iter_parts():
            if hasattr(part, "element") and str(part.partname).startswith("/word/"):
                texts += [t.text or "" for t in part.element.iter(qn("w:t"))]
        full = "".join(texts)
        chk = {"批註換掉": "王小明" not in full and "B223456789" not in full and "〔身分證" in full,
               "頁首換掉": "頁首 〔個案01〕" in full,
               "作者清空": not o.core_properties.author, "標題清空": not o.core_properties.title}
        for k, v in chk.items(): print(f"[docx {k}] {'對' if v else '錯'}"); ok &= v
        # xlsx：日期型出生日期欄、屬性；還原回日期
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["姓名", "出生日期", "受理日期"]); ws.append(["王小明", datetime.datetime(1950, 3, 2), datetime.datetime(2026, 1, 5)])
        wb.properties.creator = "王小明"; wb.save(d / "T.xlsx")
        out, mp = deid(d / "T.xlsx", names, d)
        o = openpyxl.load_workbook(out)
        chk = {"生日日期格換代號": o.active["B2"].value == "〔生日01〕", "受理日期不動": o.active["C2"].value == datetime.datetime(2026, 1, 5),
               "creator 清空": not o.properties.creator}
        for k, v in chk.items(): print(f"[xlsx {k}] {'對' if v else '錯'}"); ok &= v
        subprocess.run([PY, str(ROOT / "restore.py"), str(out), str(mp), "--out", str(d / "restored")], check=True, capture_output=True)
        r = openpyxl.load_workbook(d / "restored" / "T.xlsx").active
        v = r["B2"].value == datetime.datetime(1950, 3, 2); print(f"[xlsx 還原回日期型] {'對' if v else '錯'} {r['B2'].value!r}"); ok &= v
        # pdf：屬性
        import pymupdf
        pdf = pymupdf.open(); pg = pdf.new_page(); pg.insert_text((72, 72), "王小明 0912-345-678", fontname="china-t", fontsize=12)
        pdf.set_metadata({"title": "王小明", "author": "王小明"}); pdf.save(d / "T.pdf"); pdf.close()
        out, mp = deid(d / "T.pdf", names, d)
        o = pymupdf.open(out); meta = {k: v for k, v in o.metadata.items() if v and k != "format"}
        txt = o[0].get_text()
        chk = {"屬性清空": not meta, "正文換掉": "王小明" not in txt and "〔個案01〕" in txt}
        for k, v in chk.items(): print(f"[pdf {k}] {'對' if v else '錯'} {meta if k == '屬性清空' and meta else ''}"); ok &= v
        # jpg：EXIF
        from PIL import Image
        im = Image.new("RGB", (400, 100), "white"); ex = Image.Exif(); ex[0x010e] = "A123456789"; ex[0x013b] = "0912345678"  # EXIF 這兩欄只收 ASCII
        im.save(d / "T.jpg", exif=ex.tobytes())
        assert Image.open(d / "T.jpg").getexif().get(0x013b) == "0912345678"
        out, mp = deid(d / "T.jpg", names, d)
        v = not dict(Image.open(out).getexif()); print(f"[jpg EXIF 清空] {'對' if v else '錯'}"); ok &= v
    print("RESULT", "全對" if ok else "有錯")
    return ok


if __name__ == "__main__":
    sys.exit(0 if run() else 1)
