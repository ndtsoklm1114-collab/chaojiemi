"""xlsx：公式儲存格不被動、文字儲存格換代號、還原逐格相同。全部用假資料。"""
import sys, pathlib, tempfile, subprocess, json
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]


def run():
    with tempfile.TemporaryDirectory() as d:
        d = pathlib.Path(d)
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "名單"
        ws.append(["姓名", "身分證", "電話", "金額"])
        ws.append(["王小明", "A123456789", "0912-345-678", 100])
        ws.append(["陳大華", "B223456789", "08-7000000", 200])
        ws["D4"] = "=SUM(D2:D3)"; ws["E2"] = '=A2&"先生"'; ws["A5"] = "=COUNTA(A2:A3)"
        wb.save(d / "F.xlsx")
        (d / "names.txt").write_text("[個案]\n王小明\n陳大華\n", encoding="utf-8")
        subprocess.run([sys.executable, str(ROOT / "deidentify.py"), str(d / "F.xlsx"), "--names", str(d / "names.txt"),
                        "--out", str(d / "out"), "--map", str(d / "map")], check=True, capture_output=True)
        o = openpyxl.load_workbook(d / "out" / "F.xlsx")["名單"]
        ok = True
        for cell, want in (("D4", "=SUM(D2:D3)"), ("E2", '=A2&"先生"'), ("A5", "=COUNTA(A2:A3)")):
            print(f"[公式保留] {'對' if o[cell].value == want else '錯'} {cell} = {o[cell].value}")
            ok &= o[cell].value == want
        for cell, code in (("A2", "〔個案01〕"), ("B2", "〔身分證01〕"), ("C2", "〔電話01〕")):
            print(f"[換代號] {'對' if o[cell].value == code else '錯'} {cell} = {o[cell].value}")
            ok &= o[cell].value == code
        subprocess.run([sys.executable, str(ROOT / "restore.py"), str(d / "out" / "F.xlsx"), str(d / "map" / "F.xlsx.對照表.json"),
                        "--out", str(d / "restored")], check=True, capture_output=True)
        a, b = openpyxl.load_workbook(d / "F.xlsx"), openpyxl.load_workbook(d / "restored" / "F.xlsx")
        diff = sum(1 for ws in a.worksheets for row in ws.iter_rows() for c in row if b[ws.title][c.coordinate].value != c.value)
        print(f"[還原逐格] {'對' if diff == 0 else '錯'} 差異 {diff} 格")
        ok &= diff == 0
        print("RESULT", "全對" if ok else "有錯")
        return ok


if __name__ == "__main__":
    sys.exit(0 if run() else 1)
