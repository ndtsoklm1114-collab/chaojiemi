"""部分保留（只給 xlsx 清冊）：留前幾碼、代號只代表被遮的那段、還原接得回完整的值；沒開時行為跟以前一模一樣。全部假資料。"""
import sys, pathlib, tempfile
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from deid.core import NameList, Mapper, deidentify_text, restore_text, keep_for, DEFAULT_RETENTION
from deid.run import deidentify_file, restore_file

TEXT = "王小明 A123456789 0912-345-678 屏東縣潮州鎮四維路100號2樓 08-7891234 民國45年3月2日 內埔鄉老街12號"


def run() -> bool:
    checks = []
    nl = NameList(names={"王小明": "個案"}, skip=[])
    m = Mapper()
    out = deidentify_text(TEXT, nl, m, "A1", DEFAULT_RETENTION)
    checks.append(("身分證留前 2 碼", "A1〔身分證01〕" in out and "A123456789" not in out))
    checks.append(("手機留前 4 碼", "0912〔電話01〕" in out))
    checks.append(("市話留前 4 碼（含分隔）", "08-78〔電話02〕" in out))
    checks.append(("地址留到縣市＋鄉鎮", "屏東縣潮州鎮〔地址01〕" in out and "四維路" not in out))
    checks.append(("沒有縣市的地址留到鄉", "內埔鄉〔地址02〕" in out))
    checks.append(("姓名全遮", "〔個案01〕" in out and "王" not in out.replace("〔個案01〕", "")))
    checks.append(("生日全遮", "〔生日01〕" in out and "45年" not in out))
    entries = m.to_list()
    idv = next(e for e in entries if e["kind"] == "身分證")
    checks.append(("對照表 value 只記被遮的那段、另記保留與原值", idv["value"] == "23456789" and idv["保留"] == "A1" and idv["原值"] == "A123456789"))
    checks.append(("還原接回完整原文", restore_text(out, {"對照": entries}) == TEXT))
    checks.append(("留的字數不會等於整個值", keep_for("身分證", "A1", DEFAULT_RETENTION) == 0))
    # 沒開：跟以前一樣全遮
    m2 = Mapper()
    out2 = deidentify_text(TEXT, nl, m2)
    checks.append(("沒開部分保留時整個值換掉", "〔身分證01〕" in out2 and "A1〔" not in out2 and "保留" not in m2.to_list()[0]))
    # 整檔：xlsx 開著 → 部分保留；docx 給了 retention 也不理
    with tempfile.TemporaryDirectory() as d:
        import openpyxl, docx
        d = pathlib.Path(d)
        wb = openpyxl.Workbook(); ws = wb.active; ws.append(["姓名", "身分證", "地址"]); ws.append(["王小明", "A123456789", "屏東縣潮州鎮四維路100號2樓"]); wb.save(d / "清冊.xlsx")
        r = deidentify_file(d / "清冊.xlsx", nl, d / "out", d / "map", retention=DEFAULT_RETENTION)
        wb2 = openpyxl.load_workbook(str(r.output)); v = [c.value for c in wb2.active[2]]
        checks.append(("xlsx 開著：儲存格是 A1〔身分證01〕", v[1] == "A1〔身分證01〕" and v[2] == "屏東縣潮州鎮〔地址01〕"))
        checks.append(("xlsx 殘留檢查過（完整原值找不到）", r.residue == []))
        # 備註欄有另一種寫法的證號：殘留要抓到身分證，但不能因為「345678」碰巧在裡面就把電話也算殘留
        ws2 = openpyxl.load_workbook(d / "清冊.xlsx"); w = ws2.active; w.append(["備註", "證號 A123 456 789", "電話 0912-345-678 的被遮段 345678 出現在 A123456789 裡"]); ws2.save(d / "清冊2.xlsx")
        r3 = deidentify_file(d / "清冊2.xlsx", nl, d / "out", d / "map", retention=DEFAULT_RETENTION)
        checks.append(("另一種寫法的證號算殘留", "身分證01" in r3.residue))
        checks.append(("被遮段碰巧出現在別的數字裡不算殘留", not any(c.startswith("電話") for c in r3.residue)))
        rr = restore_file(r.output, r.mapping, d / "back")
        wb3 = openpyxl.load_workbook(str(rr.output)); v3 = [c.value for c in wb3.active[2]]
        checks.append(("xlsx 還原逐格相同", v3 == ["王小明", "A123456789", "屏東縣潮州鎮四維路100號2樓"]))
        doc = docx.Document(); doc.add_paragraph("身分證 A123456789"); doc.save(d / "報告.docx")
        r2 = deidentify_file(d / "報告.docx", nl, d / "out", d / "map", retention=DEFAULT_RETENTION)
        checks.append(("docx 給了 retention 也全遮", "A1〔" not in docx.Document(str(r2.output)).paragraphs[0].text))
    ok = True
    for name, passed in checks:
        print(("✓ " if passed else "✗ ") + name)
        ok = ok and passed
    print("RESULT", "全對" if ok else "有錯")
    return ok


if __name__ == "__main__":
    sys.exit(0 if run() else 1)
